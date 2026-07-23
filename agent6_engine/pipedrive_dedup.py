"""Pipedrive-Dedup — prueft ZWEI Dimensionen: Firma UND Anwendung.

Warum zwei Dimensionen: Eine Firma kann laengst im CRM stehen (z.B. Cormidi als
Mini-Dumper-Lead). Finden wir dieselbe Firma mit einer ANDEREN Anwendung
(z.B. Spritzbeton-Roboter), ist das eine NEUE Chance beim bestehenden Kunden —
kein Duplikat. Reines Firmen-Matching wuerde diese Chance faelschlich wegwerfen.

Ergebnis je Kandidat (Spalte 'dedup_status'):
  duplikat        -> Firma UND Anwendung/Modell bereits im CRM  -> nicht nochmal anlegen
  neue_anwendung  -> Firma im CRM, Anwendung/Modell aber NEU    -> Chance bei Bestandskunde!
  pruefen         -> Firma aehnlich (Tippfehler/Variante)       -> kurz manuell sichten
  neu             -> Firma unbekannt                            -> neuer Lead fuer SDR

Firmen-Matching nutzt die bewaehrte Fuzzy-Kaskade der OEM-Engine (Normalisierung ->
Teilstring -> Marken-Token -> Zeichen-Aehnlichkeit gegen Tippfehler, plus Whitelist),
weil exakte Vergleiche an Ein-Zeichen-Fehlern scheitern (Lehre aus dem Niftylift-Fall).
"""
from __future__ import annotations
import csv, json, argparse, re
from difflib import SequenceMatcher
from .entity_resolution import norm, brand_tokens, _whitelisted, _jaccard

MIN_SUBSTRING = 5
TOKEN_MIN_JACCARD = 0.34
TYPO_MIN_RATIO = 0.85          # Zeichen-Aehnlichkeit je Marken-Token (Tippfehler)
APP_MIN_RATIO = 0.80           # Anwendung/Modell gilt als "dieselbe"


# ---------------------------------------------------------------- Firmen-Matching
def _firma_match(name, pd_index):
    """-> (level, matched_org, score). level: exakt | aehnlich | keiner."""
    n = norm(name)
    if not n:
        return "keiner", "", 0.0
    if n in pd_index["exact"]:
        return "exakt", pd_index["exact"][n], 1.0
    for pn, orig in pd_index["exact"].items():
        if len(n) >= MIN_SUBSTRING and len(pn) >= MIN_SUBSTRING and (n in pn or pn in n):
            if not _whitelisted(name, orig):
                return "exakt", orig, 0.9
    tk = brand_tokens(name)
    best, best_score = "", 0.0
    for orig, ptk, _apps in pd_index["orgs"]:
        j = _jaccard(tk, ptk)
        if j > best_score and not _whitelisted(name, orig):
            best, best_score = orig, j
    if best_score >= TOKEN_MIN_JACCARD:
        return "exakt", best, round(best_score, 2)
    # Tippfehler-Stufe
    tbest, tscore = "", 0.0
    for orig, ptk, _apps in pd_index["orgs"]:
        for a in tk:
            for b in ptk:
                if abs(len(a) - len(b)) > 2:
                    continue
                r = SequenceMatcher(None, a, b).ratio()
                if r > tscore and not _whitelisted(name, orig):
                    tbest, tscore = orig, r
    if tscore >= TYPO_MIN_RATIO:
        return "aehnlich", tbest, round(tscore, 2)
    return "keiner", "", 0.0


# ---------------------------------------------------------------- Anwendungs-Matching
def _norm_app(s) -> str:
    """Normalisiert Anwendung/Modell fuer den Vergleich."""
    t = norm(s)
    t = re.sub(r"[^a-z0-9äöüß ]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def _app_bekannt(anwendung, modell, bekannte_apps) -> tuple[bool, str]:
    """Ist diese Anwendung/dieses Modell bei der Firma schon im CRM?"""
    kand = [_norm_app(x) for x in (anwendung, modell) if str(x or "").strip()]
    if not kand or not bekannte_apps:
        return False, ""
    for k in kand:
        for a in bekannte_apps:
            na = _norm_app(a)
            if not na or not k:
                continue
            if k == na or (len(k) >= 4 and len(na) >= 4 and (k in na or na in k)):
                return True, a
            if SequenceMatcher(None, k, na).ratio() >= APP_MIN_RATIO:
                return True, a
    return False, ""


# ---------------------------------------------------------------- Index & Status
def build_index(pd_rows) -> dict:
    """pd_rows: Liste von dicts {'name':..., 'anwendung':...} oder reine Namens-Strings."""
    exact, orgs = {}, []
    per_org = {}
    for row in pd_rows:
        if isinstance(row, str):
            name, app = row, ""
        else:
            name = row.get("name", "")
            app = row.get("anwendung", "")
        name = str(name or "").strip()
        if not name:
            continue
        exact[norm(name)] = name
        per_org.setdefault(name, set())
        if app:
            per_org[name].add(str(app).strip())
    for name, apps in per_org.items():
        orgs.append((name, brand_tokens(name), apps))
    return {"exact": exact, "orgs": orgs}


def match_status(name, anwendung, modell, pd_index) -> tuple[str, str, str, float]:
    """-> (status, matched_org, matched_app, score)."""
    level, org, score = _firma_match(name, pd_index)
    if level == "keiner":
        return "neu", "", "", 0.0
    apps = next((a for (o, _t, a) in pd_index["orgs"] if o == org), set())
    bekannt, matched_app = _app_bekannt(anwendung, modell, apps)
    if level == "aehnlich":
        return "pruefen", org, matched_app, score
    if bekannt:
        return "duplikat", org, matched_app, score
    return "neue_anwendung", org, "", score


# ---------------------------------------------------------------- IO
def _pick_col(cols, cands):
    """Erst exakte Uebereinstimmung, dann Teilstring — Pipedrive exportiert
    Spalten wie 'Organisation - Name' oder 'Deal - Titel'."""
    low = [(c, str(c).strip().lower()) for c in cols if c]
    for c, lc in low:
        if lc in cands:
            return c
    for c, lc in low:
        for cand in cands:
            if cand in lc:
                return c
    return None


_NAME_COLS = ("organisation", "organization", "firma", "company", "firmenname", "kunde",
              "org name", "organization name", "organisation name", "name")
_APP_COLS = ("anwendung", "application", "deal", "deal title", "titel", "title",
             "produkt", "product", "maschine", "machine", "modell", "model",
             "deal - title", "segment", "beschreibung")


def load_pipedrive(path) -> list:
    """CSV/XLSX/JSON mit Firmen + (optional) Anwendung/Deal-Titel.
    Erkannte Spalten: Name/Organization/Firma  und  Anwendung/Deal/Titel/Produkt/Maschine."""
    p = str(path).lower()
    if p.endswith(".json"):
        data = json.load(open(path, encoding="utf-8"))
        out = []
        for d in data:
            if isinstance(d, dict):
                out.append({"name": d.get("name", ""), "anwendung": d.get("anwendung", d.get("title", ""))})
            else:
                out.append({"name": str(d), "anwendung": ""})
        return out
    if p.endswith((".xlsx", ".xlsm", ".xltx")):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)
        # Kopfzeile suchen (erste Zeile mit einer erkannten Namensspalte, max. 5 Zeilen)
        header, header_row = None, None
        for i, r in enumerate(rows_iter):
            if not r:
                continue
            cells = [str(c).strip() if c is not None else "" for c in r]
            if _pick_col(cells, _NAME_COLS):
                header, header_row = cells, i
                break
            if i >= 5:
                break
        if header is None:
            wb2 = load_workbook(path, read_only=True, data_only=True)
            ws2 = wb2[wb2.sheetnames[0]]
            first = next(ws2.iter_rows(values_only=True), ())
            print("  ⚠ Keine Firmen-Spalte erkannt. Gefundene Spalten in Zeile 1:",
                  [str(c) for c in first if c is not None][:15])
            return []
        c_name = _pick_col(header, _NAME_COLS)
        c_app = _pick_col(header, _APP_COLS)
        print(f"  Excel-Spalten erkannt -> Firma: '{c_name}' | Anwendung: '{c_app}'")
        i_name = header.index(c_name)
        i_app = header.index(c_app) if c_app else None
        rows = []
        for r in ws.iter_rows(min_row=header_row + 2, values_only=True):
            if not r or i_name >= len(r) or not r[i_name]:
                continue
            rows.append({"name": str(r[i_name]).strip(),
                         "anwendung": str(r[i_app]).strip() if (i_app is not None and i_app < len(r) and r[i_app]) else ""})
        return rows
    rows = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        rd = csv.DictReader(f)
        cols = rd.fieldnames or []
        c_name = _pick_col(cols, _NAME_COLS)
        c_app = _pick_col(cols, _APP_COLS)
        for row in rd:
            rows.append({"name": row.get(c_name, "") if c_name else "",
                         "anwendung": row.get(c_app, "") if c_app else ""})
    return rows


def run(candidates_path, pipedrive_path, out_csv="agent6_dedup.csv", airtable=None) -> dict:
    if str(candidates_path).lower().endswith(".json"):
        cands = json.load(open(candidates_path, encoding="utf-8"))
    else:
        cands = list(csv.DictReader(open(candidates_path, newline="", encoding="utf-8-sig")))
    pd_rows = load_pipedrive(pipedrive_path)
    idx = build_index(pd_rows)
    rows, stats = [], {"kandidaten": len(cands), "pipedrive_zeilen": len(pd_rows),
                       "pipedrive_firmen": len(idx["orgs"]),
                       "duplikat": 0, "neue_anwendung": 0, "pruefen": 0, "neu": 0}
    for c in cands:
        name = c.get("oem") or c.get("oem_guess") or ""
        anwendung = c.get("anwendung", "")
        modell = c.get("modell", "")
        status, org, app, score = match_status(name, anwendung, modell, idx)
        stats[status] += 1
        r = dict(c)
        r.update({"dedup_status": status, "pipedrive_firma": org,
                  "pipedrive_anwendung": app, "match_score": score})
        rows.append(r)
    if rows:
        with open(out_csv, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
            w.writeheader()
            for r in rows:
                w.writerow(r)
    if airtable and getattr(airtable, "enabled", False):
        stats["airtable_updated"] = airtable.update_fields(
            rows, {"Dedup-Status": "dedup_status", "Pipedrive-Firma": "pipedrive_firma"})
    stats["out_csv"] = out_csv
    return stats


def main():
    ap = argparse.ArgumentParser(description="Pipedrive-Dedup: Firma UND Anwendung pruefen")
    ap.add_argument("--candidates", default="agent6_enriched.csv")
    ap.add_argument("--pipedrive", required=True,
                    help="CSV/XLSX/JSON-Export aus Pipedrive (Spalten: Name + Anwendung/Deal/Modell)")
    ap.add_argument("--out", default="agent6_dedup.csv")
    ap.add_argument("--airtable", action="store_true",
                    help="Dedup-Status zurueck nach Airtable schreiben")
    args = ap.parse_args()
    at = None
    if args.airtable:
        from .image_search_harvester import AirtableWriter
        at = AirtableWriter()
    res = run(args.candidates, args.pipedrive, args.out, airtable=at)
    print("PIPEDRIVE-DEDUP:", res)
    print(f"  neu ({res['neu']}) + neue_anwendung ({res['neue_anwendung']}) -> an SDR")
    print(f"  pruefen ({res['pruefen']}) -> kurz sichten | duplikat ({res['duplikat']}) -> nichts tun")


if __name__ == "__main__":
    main()
