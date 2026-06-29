"""DB -> Master-Export. Erzeugt eine Export-Sicht im Hauptsheet-Format
(gleiche Spaltenüberschriften wie der Master), gespeist aus der DB.
Dies ist die generierte Sicht — die DB ist die Wahrheitsquelle.
"""
from __future__ import annotations
import datetime
from openpyxl import Workbook
from sqlalchemy import select
from .db import Organization, Claim, Model

HEADERS = ["OEM / Hersteller", "Segment", "Modell / Serie", "Land", "Domain",
           "Status", "Pattern", "LEI", "VAT", "Adresse", "Confidence",
           "Zuletzt verifiziert", "Zeile", "Quellen (Claims)"]


def export_merge_review(session, out_path: str) -> str:
    """Exportiert die nicht-auto-gemergten Kandidaten zur Mensch-Review."""
    from .db import MergeCandidate
    wb = Workbook()
    ws = wb.active
    ws.title = "Merge-Review"
    ws.append(["#", "Neuer OEM", "Möglicher Treffer", "Tier", "Score", "Grund", "Entscheidung"])
    for i, mc in enumerate(session.scalars(select(MergeCandidate).order_by(MergeCandidate.score.desc())), start=1):
        ws.append([i, mc.new_name, mc.matched_name, mc.tier, round(mc.score, 2), mc.reason, mc.decision])
    wb.save(out_path)
    return out_path


def export_master(session, out_path: str, header_row: int = 2,
                  sheet_name: str = "Agent 6 — Crawler OEM") -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.cell(1, 1, f"Agent 6 — Export {datetime.date.today().isoformat()}")
    for j, h in enumerate(HEADERS, start=1):
        ws.cell(header_row, j, h)

    r = header_row
    for i, org in enumerate(session.scalars(select(Organization).order_by(Organization.id)), start=1):
        r += 1
        model = session.scalar(select(Model).where(Model.organization_id == org.id))
        pattern = model.pattern_code if model else None
        segment = model.segment if model else None
        claims = list(session.scalars(select(Claim).where(
            Claim.subject_type == "organization", Claim.subject_id == org.id)))
        srcs = "; ".join(f"{c.predicate}={c.object_value}[{c.source_type}]" for c in claims[:6])
        row = [org.canonical_name, segment, (model.name if model else None), org.country, org.domain,
               org.status, pattern, org.lei, org.vat, org.address, round(org.confidence, 2),
               org.last_verified.isoformat() if org.last_verified else None, i, srcs]
        for j, v in enumerate(row, start=1):
            ws.cell(r, j, v)
    wb.save(out_path)
    return out_path
