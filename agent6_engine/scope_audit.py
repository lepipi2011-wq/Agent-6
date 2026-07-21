"""Scope-Audit — pro OEM ein Urteil, ob er überhaupt in-scope ist.

Gruppiert die bereits geernteten Bilder nach OEM (Dateiname <Class>__<OEM>__<idx>),
schickt je OEM eine Stichprobe an den VLM und fragt: baut diese Firma eine
kettengetriebene (Raupen-)Maschine mit Fernsteuerung/kabelgebundenem HMI?
Ergebnis: eine Review-xlsx (OEM | Klasse | #Bilder | Verdikt | Begründung) —
die abarbeitbare Drop-Liste, die die falschen OEMs aus dem Master entfernt.

Live braucht ANTHROPIC_API_KEY. Gruppierung + Parsing sind offline testbar.
"""
from __future__ import annotations
import os, io, json, base64, argparse, glob
from collections import defaultdict

_MEDIA = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "webp": "image/webp"}


def group_images_by_oem(root: str) -> dict:
    """{ (class, oem): [pfade...] } aus Dateinamen <Class>__<OEM>__<idx>.ext."""
    groups = defaultdict(list)
    for path in glob.glob(os.path.join(root, "**", "*.*"), recursive=True):
        base = os.path.basename(path)
        if base.rsplit(".", 1)[-1].lower() not in _MEDIA:
            continue
        parts = base.rsplit(".", 1)[0].split("__")
        if len(parts) >= 3:
            cls, oem = parts[0], parts[1]
            groups[(cls, oem.replace("-", " "))].append(path)
    return groups


def _b64(path):
    with open(path, "rb") as f:
        raw = f.read()
    return base64.standard_b64encode(raw).decode(), _MEDIA.get(path.rsplit(".", 1)[-1].lower(), "image/jpeg")


def _prep_image(path, max_px=1568):
    """Verkleinert auf API-taugliche Größe (max_px lange Kante, JPEG). Behebt
    BadRequestError bei Riesenbildern (Decompression-Bomb). Gibt (b64, media)."""
    from PIL import Image
    im = Image.open(path).convert("RGB")
    w, h = im.size
    if max(w, h) > max_px:
        s = max_px / max(w, h)
        im = im.resize((max(1, int(w * s)), max(1, int(h * s))))
    buf = io.BytesIO()
    im.save(buf, format="JPEG", quality=85)
    return base64.standard_b64encode(buf.getvalue()).decode(), "image/jpeg"


def parse_audit_verdict(txt: str) -> tuple[str, str]:
    """-> (verdict in {keep,drop,unsure}, reason). Offline testbar."""
    s = (txt or "").strip()
    a, b = s.find("{"), s.rfind("}")
    if a >= 0 and b > a:
        try:
            d = json.loads(s[a:b + 1])
            v = str(d.get("verdict", "unsure")).lower()
            if v not in ("keep", "drop", "unsure"):
                v = "unsure"
            return v, str(d.get("reason", ""))[:120]
        except Exception:
            pass
    low = s.lower()
    if "drop" in low:
        return "drop", "geparst:drop"
    if "keep" in low:
        return "keep", "geparst:keep"
    return "unsure", "unklar"


def vlm_auditor(cfg):
    """Default-Auditor: nutzt Claude multimodal. Kein Key -> ('unsure','kein-key')."""
    def _audit(oem, cls, paths):
        key = os.environ.get("ANTHROPIC_API_KEY")
        if not key:
            return "unsure", "kein-key"
        try:
            import anthropic
            client = anthropic.Anthropic(api_key=key)
            model = cfg["harvest"].get("vlm_model", "claude-haiku-4-5-20251001")
            content = []
            for p in paths[:cfg.get("audit_sample", 3)]:
                try:
                    data, media = _prep_image(p)
                except Exception:
                    continue                      # unlesbares Bild überspringen
                content.append({"type": "image", "source": {"type": "base64",
                                "media_type": media, "data": data}})
            frage = (
                f"Firmenname: {oem}. Erwartetes Pattern: {cls}.\n"
                "Baut diese Firma eine KETTENGETRIEBENE (Raupen-)Maschine mit "
                "Fernsteuerung/kabelgebundenem HMI (Safety-Remote-Control-Kandidat)? "
                "NICHT in-scope: Radmaschinen/Radlader/Stapler, LKW, handgeführte Geräte/"
                "Messtechnik/Sonden, Schiffe/Marine, Krane/Hubarbeitsbühnen ohne Kette, "
                "Händler/Distributoren, reine Software/Komponenten. "
                'Antworte NUR JSON {"verdict":"keep|drop|unsure","reason":"kurz"}.')
            if not content:
                # Kein Bild vorhanden -> Urteil aus Name/Segment, klar gekennzeichnet
                frage += ("\nHINWEIS: Es liegen KEINE Bilder vor. Urteile nur nach Firmenname "
                          "und Segment; bei echter Unsicherheit verdict='unsure'.")
                content = [{"type": "text", "text": frage}]
            else:
                content.append({"type": "text", "text": frage})
            msg = client.messages.create(model=model, max_tokens=150,
                                         messages=[{"role": "user", "content": content}])
            txt = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")
            return parse_audit_verdict(txt)
        except Exception as e:
            return "unsure", f"vlm-fehler:{type(e).__name__}"
    return _audit


def load_master_scopes(cfg) -> dict:
    """{normalisierter OEM-Name: scope} aus dem Master (Spalte Label/Scope)."""
    try:
        from . import master_io as M
        wb = M.open_master(cfg)
        seeds = M.read_seeds(wb, cfg, M.read_patterns(wb, cfg))
        return {s.oem.strip().lower(): getattr(s, "scope", "") for s in seeds if s.oem}
    except Exception:
        return {}


def load_master_seeds(cfg) -> dict:
    """{norm. OEM: Seed} — damit auch OEMs OHNE geerntete Bilder auditiert werden können."""
    try:
        from . import master_io as M
        wb = M.open_master(cfg)
        seeds = M.read_seeds(wb, cfg, M.read_patterns(wb, cfg))
        return {s.oem.strip().lower(): s for s in seeds if s.oem}
    except Exception:
        return {}


def run_audit(cfg, root=None, out_xlsx="agent6_scope_audit.xlsx", limit=None, auditor=None,
              use_master=True, only_open=False, include_imageless=False) -> dict:
    """only_open=True  -> nur OEMs mit Scope 'open' oder ungelabelt (spart Kosten,
                          respektiert bereits getroffene Entscheidungen).
    include_imageless=True -> auch OEMs ohne geerntete Bilder (Urteil nur aus Name+Segment)."""
    root = root or os.path.join(cfg["output"]["dir"], "raw")
    if not os.path.isdir(root):
        root = cfg["output"]["dir"]
    audit = auditor if auditor is not None else vlm_auditor(cfg)
    groups = group_images_by_oem(root)
    scopes = load_master_scopes(cfg) if (use_master or only_open) else {}
    seeds = load_master_seeds(cfg) if include_imageless else {}

    # Arbeitsliste: OEMs mit Bildern + (optional) OEMs ohne Bilder aus dem Master
    work = {(cls, oem): paths for (cls, oem), paths in groups.items()}
    if include_imageless:
        have = {oem.strip().lower() for (_c, oem) in work}
        for key, s in seeds.items():
            if key not in have:
                work[((s.pattern or "?"), s.oem)] = []

    rows, stats = [], {"oems": 0, "keep": 0, "drop": 0, "unsure": 0,
                       "aus_master": 0, "skipped_decided": 0, "ohne_bilder": 0}
    for i, ((cls, oem), paths) in enumerate(sorted(work.items())):
        pre = scopes.get(oem.strip().lower(), "")
        if only_open and pre in ("exclude", "lead"):
            stats["skipped_decided"] += 1      # bereits entschieden -> kein VLM-Call
            continue
        if limit and stats["oems"] >= limit:
            break
        if not only_open and use_master and pre == "exclude":
            verdict, reason = "drop", "aus Master (Ausschluss/Adjacent/Distributor)"
            stats["aus_master"] += 1
        elif not only_open and use_master and pre == "lead":
            verdict, reason = "keep", "aus Master (Crawler-Lead)"
            stats["aus_master"] += 1
        else:
            verdict, reason = audit(oem, cls, paths)
            if not paths:
                stats["ohne_bilder"] += 1
        stats["oems"] += 1
        stats[verdict] = stats.get(verdict, 0) + 1
        rows.append((oem, cls, len(paths), verdict, reason,
                     "; ".join(os.path.basename(p) for p in paths[:3]) or "(keine Bilder)"))
    _write_xlsx(rows, out_xlsx)
    stats["out"] = out_xlsx
    return stats


def _write_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font
    wb = Workbook(); ws = wb.active; ws.title = "Scope-Audit"
    ws.append(["OEM", "Klasse", "#Bilder", "Verdikt", "Begründung", "Beispielbilder"])
    for c in ws[1]:
        c.font = Font(bold=True)
    red = PatternFill("solid", fgColor="FFC7CE"); green = PatternFill("solid", fgColor="C6EFCE")
    for r in rows:
        ws.append(list(r))
        fill = red if r[3] == "drop" else (green if r[3] == "keep" else None)
        if fill:
            ws.cell(ws.max_row, 4).fill = fill
    for col, w in zip("ABCDEF", (34, 10, 9, 10, 60, 50)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def main():
    from .config import load_config
    ap = argparse.ArgumentParser(description="Scope-Audit: OEMs auf In-Scope prüfen")
    ap.add_argument("--config", default="config.yaml")
    ap.add_argument("--root", default=None, help="Bilder-Wurzel (default: output.dir/raw)")
    ap.add_argument("--out", default="agent6_scope_audit.xlsx")
    ap.add_argument("--sample", type=int, default=3, help="Bilder pro OEM ans VLM")
    ap.add_argument("--limit", type=int, default=None)
    ap.add_argument("--only-open", dest="only_open", action="store_true",
                    help="nur OEMs mit Scope 'offen'/ungelabelt (spart Kosten)")
    ap.add_argument("--include-imageless", dest="include_imageless", action="store_true",
                    help="auch OEMs ohne geerntete Bilder (Urteil aus Name/Segment)")
    args = ap.parse_args()
    cfg = load_config(args.config); cfg["audit_sample"] = args.sample
    res = run_audit(cfg, root=args.root, out_xlsx=args.out, limit=args.limit,
                    only_open=args.only_open, include_imageless=args.include_imageless)
    print("SCOPE-AUDIT:", res)


if __name__ == "__main__":
    main()
