"""Master-I/O — header-basiert, struktur-erhaltend.

Liest die Saatliste (bestätigte OEMs), die Pattern-DB und gibt beides
als einfache Datenstrukturen zurück. Schreibt NICHT in den Master.
"""
from __future__ import annotations
import re
from dataclasses import dataclass, field
from openpyxl import load_workbook


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def _col_map(ws, header_row):
    m = {}
    for j in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=j).value
        if v is not None:
            m[_norm(v)] = j
    return m


def _resolve(cmap, *subs):
    for hdr, j in cmap.items():
        if all(s in hdr for s in subs):
            return j
    return None


def _domain(url):
    if not url:
        return None
    m = re.match(r"https?://([^/]+)", str(url).strip())
    if not m:
        return None
    host = m.group(1)
    return f"https://{host}"


def _pattern_from_segment(seg):
    m = re.match(r"\s*(P\d+)\b", str(seg or ""))
    return m.group(1) if m else None


_EXCL_KEYWORDS = ("adjacent", "exclusion", "ausschluss", "anti-pattern", "anti pattern",
                  "agv", "spmt", "schreitbagger", "walking excavator", "spider excavator",
                  "tbm", "tunnel boring", "sewer", "kanalinspektion", "pipe inspection",
                  "wheeled", "radgetrieben", "radlader")


def normalize_class(raw):
    """Rohes Pattern-Label -> saubere Trainingsklasse.
    '^P\\d+' -> P-Code; Adjacent/Exklusion/Anti-Pattern -> 'out_scope'; leer/unbekannt -> 'unknown'.
    Verhindert Splitter-Klassen aus der Pattern-Spalte des reconciled-Masters."""
    s = str(raw or "").strip()
    if not s:
        return "unknown"
    if s.lower() in ("out_scope", "unknown", "in_scope"):
        return s.lower()
    m = re.match(r"^\s*(P\d+)\b", s)
    if m:
        return m.group(1)
    low = s.lower()
    if any(k in low for k in _EXCL_KEYWORDS):
        return "out_scope"
    return "unknown"


@dataclass
class Seed:
    oem: str
    pattern: str | None
    segment: str | None
    land: str | None
    domain: str | None
    link: str | None
    row: int


@dataclass
class Pattern:
    code: str
    img_query: str | None = None
    text_terms: str | None = None
    ref_oem: str | None = None
    segment: str | None = None
    candidates: str | None = None


def read_patterns(wb, cfg) -> dict[str, Pattern]:
    name = cfg["master"]["pattern_sheet"]
    if name not in wb.sheetnames:
        return {}
    ws = wb[name]
    hr = cfg["master"]["pattern_header_row"]
    pm = _col_map(ws, hr)
    c_code = _resolve(pm, "pattern", "code") or 2
    c_img = _resolve(pm, "bildsuch") or 5
    c_ref = _resolve(pm, "referenz-oem") or 6
    c_seg = _resolve(pm, "segment") or 9
    c_cand = _resolve(pm, "bekannte kandidaten") or 12
    c_text = _resolve(pm, "textsuch") or 13
    out = {}
    for r in range(hr + 1, ws.max_row + 1):
        code = ws.cell(r, c_code).value
        if not code or not re.match(r"^P\d+", str(code).strip()):
            continue
        code = str(code).strip().split()[0]
        out[code] = Pattern(
            code=code,
            img_query=ws.cell(r, c_img).value,
            text_terms=ws.cell(r, c_text).value,
            ref_oem=ws.cell(r, c_ref).value,
            segment=ws.cell(r, c_seg).value,
            candidates=ws.cell(r, c_cand).value,
        )
    return out


def read_seeds(wb, cfg, patterns: dict[str, Pattern]) -> list[Seed]:
    ws = wb[cfg["master"]["main_sheet"]]
    hr = cfg["master"]["header_row"]
    cm = _col_map(ws, hr)
    c_oem = _resolve(cm, "oem") or _resolve(cm, "hersteller")
    c_seg = _resolve(cm, "segment")
    c_link = _resolve(cm, "link")
    c_land = _resolve(cm, "land")
    c_pat = _resolve(cm, "pattern")  # in reconciled-Version evtl. vorhanden

    # OEM-Name -> Pattern aus Pattern-DB-Kandidaten (Fallback-Lookup)
    cand_index = {}
    for p in patterns.values():
        for tok in re.split(r"[;,/]", str(p.candidates or "")):
            tok = tok.strip()
            if len(tok) >= 4:
                cand_index[_norm(tok)] = p.code

    seeds = []
    for r in range(hr + 1, ws.max_row + 1):
        oem = ws.cell(r, c_oem).value
        if not oem or not str(oem).strip():
            continue
        seg = ws.cell(r, c_seg).value if c_seg else None
        link = ws.cell(r, c_link).value if c_link else None
        land = ws.cell(r, c_land).value if c_land else None
        pat = (ws.cell(r, c_pat).value if c_pat else None) or _pattern_from_segment(seg)
        if not pat:
            non = _norm(oem)
            for key, code in cand_index.items():
                if key in non or non in key:
                    pat = code
                    break
        seeds.append(Seed(oem=str(oem).strip(),
                          pattern=(str(pat).strip() if pat else None),
                          segment=str(seg).strip() if seg else None,
                          land=str(land).strip() if land else None,
                          domain=_domain(link), link=str(link).strip() if link else None,
                          row=r))
    return seeds


def open_master(cfg):
    # read_only=False: wir nutzen wahlfreien .cell()-Zugriff; read_only macht das O(n^2).
    return load_workbook(cfg["master"]["path"], read_only=False, data_only=True)
