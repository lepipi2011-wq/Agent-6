"""Review-Rückspielung — Entscheidungen aus dem Merge-Review-xlsx ausführen.

Liest das von db_export.export_merge_review erzeugte Sheet (Spalten:
#, Neuer OEM, Möglicher Treffer, Tier, Score, Grund, Entscheidung).
Für 'merge' wird der neue OEM in den möglichen Treffer gemerged; 'keep-separate'
wird als Entscheidung vermerkt. Alles andere ('pending'/leer) bleibt offen.
"""
from __future__ import annotations
from openpyxl import load_workbook
from sqlalchemy import select
from .db import Organization, MergeCandidate
from .entity_resolution import merge_orgs, norm


def _find_org(session, name):
    n = norm(name)
    return session.scalar(select(Organization).where(Organization.norm_name == n)) \
        or session.scalar(select(Organization).where(Organization.canonical_name == str(name).strip()))


def apply_review(session, xlsx_path: str, sheet: str = "Merge-Review") -> dict:
    wb = load_workbook(xlsx_path)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    stats = {"merged": 0, "kept_separate": 0, "skipped": 0, "not_found": 0}
    for r in rows:
        if not r or len(r) < 7:
            continue
        new_name, matched_name, decision = r[1], r[2], str(r[6] or "").strip().lower()
        if decision == "merge":
            new_org = _find_org(session, new_name)
            survivor = _find_org(session, matched_name)
            if not new_org or not survivor:
                stats["not_found"] += 1
                continue
            merge_orgs(session, survivor, new_org, tier="manual-review", reason="review:merge")
            stats["merged"] += 1
        elif decision in ("keep-separate", "keep_separate", "getrennt"):
            mc = session.scalar(select(MergeCandidate).where(
                MergeCandidate.new_name == str(new_name).strip(),
                MergeCandidate.matched_name == str(matched_name).strip()))
            if mc:
                mc.decision = "keep-separate"
            stats["kept_separate"] += 1
        else:
            stats["skipped"] += 1
    session.commit()
    return stats
